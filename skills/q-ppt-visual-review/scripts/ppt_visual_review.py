#!/usr/bin/env python3
"""Heuristic PowerPoint visual review report generator."""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
import math
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile


EMU_PER_INCH = 914400
MIN_BODY_FONT_PT = 10.0
MIN_VISIBLE_FONT_PT = 9.0
EDGE_MARGIN_IN = 0.15
MIN_TEXT_CONTAINER_INSET_IN = 0.14
MIN_SHAPE_TEXT_MARGIN_IN = 0.10
MIN_NORMAL_TEXT_CONTRAST = 4.5
MIN_LARGE_TEXT_CONTRAST = 3.0
MIN_EVIDENCE_IMAGE_WIDTH_IN = 3.2
MIN_EVIDENCE_IMAGE_HEIGHT_IN = 1.6
MIN_EVIDENCE_IMAGE_AREA_IN2 = 5.0
ORPHAN_LINE_UNITS = 2.0
BADGE_CENTER_TOLERANCE_IN = 0.06
ROW_LABEL_CENTER_TOLERANCE_IN = 0.06
COMPONENT_CENTER_TOLERANCE_IN = 0.04
CIRCLE_TEXT_CENTER_TOLERANCE_IN = 0.025
RULE_LABEL_CENTER_TOLERANCE_IN = 0.04
SEPARATOR_TEXT_GAP_IN = 0.16
HEADER_ALIGN_TOLERANCE_IN = 0.10
TITLE_DRIFT_TOLERANCE_IN = 0.08
HEADER_RELATION_TOLERANCE_IN = 0.18
MIN_GRAPHIC_OVERLAP_AREA_IN2 = 0.03
GRAPHIC_OVERLAP_RATIO = 0.18
CONNECTOR_CROSSING_TOLERANCE_IN = 0.04
MIN_GRAPHIC_CONTAINER_INSET_IN = 0.08


MANUAL_CHECKLIST = [
    "Review every exported PNG at presentation size; do not accept the deck from the heuristic report alone.",
    "For generated decks, reject repeated rows, badges, circles, cards, or process steps that are built from separate per-element offsets instead of one helper-owned component contract.",
    "For editable PPTX output, prefer actual PowerPoint groups for repeated multi-shape components after the helper lays out the internal grid; grouped children must still be inspected.",
    "For generated decks, reject large foreign-color readability patches such as a white rectangle on a dark style; fix contrast through the palette, typography, and component surface instead.",
    "Check text/background contrast: ordinary required text should meet about 4.5:1 contrast, and large/bold display text should meet about 3:1. Dark text on a dark panel, or pale text on a pale panel, is a blocker.",
    "Check every card, callout, and framed text component for internal padding. Text boxes should not touch or cross the frame; generated components should reserve a helper-owned safe inset before rendering text.",
    "Check chart bars, media-slot bars, and placeholder figures as parent-child systems: child bars must stay inside parent frames with visible inset on both sides.",
    "Check note rows, bullet rows, icon-label rows, and side-panel lists: marker/circle/icon and adjacent text should be one row-center component.",
    "Reject evidence screenshots or figures whose embedded labels are unreadable at exported-slide size. Enlarge the figure, crop to the relevant region, or replace it with a simplified diagram.",
    "For process slides, reject one large shared band with connector lines floating inside it; use individual step cards/nodes and connect their boundaries or centers.",
    "For diagram slides, check non-text graphical overlaps too: icons, symbols, dashed regions, node boxes, and connector lines should not visibly cover each other unless one object is an intentional child inside a component.",
    "Check number badge + adjacent title/detail rows: the badge, title, and paired detail text should share the same visual center line.",
    "Check numeric circles: the number text frame must be the same box as the circle or centered on the circle with a middle anchor; a manually nudged text box inside an oval is a component defect.",
    "Check short colored row labels + adjacent text rows: the label and every same-row text block should share one visual center line.",
    "Check every short accent rule + adjacent label/callout text pair: they should be emitted by one helper/group and share a vertical center line.",
    "Check agenda/list blocks near template separator lines: leave visible air below the line before the first item.",
    "Check repeated top headers across the deck: section label, accent rule, and main title should follow one explicit grid or baseline contract.",
    "Check global title placement: generated decks should not have floating top rules or section labels that are visually detached from the main title.",
    "Check section divider subtitles: long subtitles must wrap or shorten before they touch the template art boundary.",
    "Check long labels and callouts: if the sentence visually wants a line break, add one instead of relying on PowerPoint auto-wrap.",
    "Check bottom callouts and small source notes for orphan punctuation, clipped descenders, or text pressed against borders.",
    "Check repeated cards, tiles, process labels, and compact rows for one punctuation convention across peer items: phrase-style usually omits terminal punctuation; sentence-style uses it consistently.",
    "After visual checks, review visible text for accidental Chinese/English mixing, leftover production notes, and wording that does not fit the audience.",
    "Confirm the slide's title, labels, and callouts express a clear technical meaning; flag vague or misleading semantics even when the layout is clean.",
    "Treat heuristic findings as leads, not final visual truth.",
]

ALLOWED_MIXED_TERMS = {
    "A346",
    "A",
    "AB",
    "ADC",
    "AI",
    "AN14856",
    "ANFIS",
    "ANN",
    "API",
    "B",
    "BOOST",
    "BRIGHT",
    "CAN",
    "DI",
    "DEMO",
    "DP",
    "DSP",
    "DV",
    "ENERGIES",
    "FLASH",
    "FPU",
    "FTO",
    "GA",
    "GMPP",
    "GWO",
    "HIL",
    "I",
    "I2C",
    "IEC",
    "INCCOND",
    "INTERNAL",
    "LLC",
    "LMPP",
    "MCU",
    "ML",
    "MPP",
    "MPPT",
    "O",
    "P",
    "POC",
    "PV",
    "PWM",
    "QU",
    "RAM",
    "ROM",
    "PSO",
    "SPI",
    "UART",
    "USB",
    "V",
}

COMMON_VISIBLE_TEXT_RISKS = [
    (re.compile(r"\bSource\s*:", re.IGNORECASE), "visible source note uses English helper wording; localize it or move it to notes if this is a polished Chinese deck"),
    (re.compile(r"\bDiscussion target\s*:", re.IGNORECASE), "visible discussion prompt uses English helper wording; localize it or move it to speaker notes"),
    (re.compile(r"\bInternal framing\b", re.IGNORECASE), "visible internal-production wording should be localized or moved to notes"),
    (re.compile(r"\bClosing slide\b", re.IGNORECASE), "visible template-production wording should not remain on a slide"),
    (re.compile(r"\bcustomer discussion deck\b", re.IGNORECASE), "visible process wording may be too internal; check audience fit"),
    (re.compile(r"\bfallback\b", re.IGNORECASE), "English `fallback` appears in Chinese context; prefer a deliberate term such as `回退` or `兜底` unless this is an interface field"),
    (re.compile(r"\bduty\b", re.IGNORECASE), "English `duty` appears in Chinese context; prefer `占空比` unless referring to a code field"),
    (re.compile(r"\bdemo\b", re.IGNORECASE), "English `demo` appears in Chinese context; prefer `演示` unless it is an official product name"),
    (re.compile(r"\btracking\b", re.IGNORECASE), "English `tracking` appears in Chinese context; prefer `跟踪` unless it is an official term"),
]


@dataclass
class Box:
    left: int
    top: int
    width: int
    height: int

    @property
    def right(self) -> int:
        return self.left + self.width

    @property
    def bottom(self) -> int:
        return self.top + self.height

    @property
    def center_x(self) -> float:
        return self.left + self.width / 2

    @property
    def center_y(self) -> float:
        return self.top + self.height / 2


@dataclass
class TextShape:
    label: str
    box: Box
    text: str
    vertical_middle: bool
    min_font_pt: float
    max_font_pt: float
    color_rgb: tuple[int, int, int] | None = None
    rotation: float = 0.0


@dataclass
class LineShape:
    label: str
    x1: int
    y1: int
    x2: int
    y2: int
    width_pt: float = 1.0


@dataclass
class VisualShape:
    label: str
    box: Box
    kind: str


@dataclass
class SlideHeaderMetrics:
    slide_idx: int
    title_label: str | None = None
    title_box: Box | None = None
    title_text: str = ""
    header_label: str | None = None
    header_box: Box | None = None
    header_text: str = ""
    rule_label: str | None = None
    rule_box: Box | None = None


def iter_child_shapes(shapes, prefix: str, top_index: int):
    for child_idx, shape in enumerate(shapes, start=1):
        label = f"{prefix}.{child_idx}"
        yield label, shape, top_index
        if hasattr(shape, "shapes"):
            yield from iter_child_shapes(shape.shapes, label, top_index)


def iter_slide_shapes(slide):
    for shape_idx, shape in enumerate(slide.shapes, start=1):
        label = f"shape {shape_idx}"
        yield label, shape, shape_idx
        if hasattr(shape, "shapes"):
            yield from iter_child_shapes(shape.shapes, label, shape_idx)


def box_overlap(a: Box, b: Box) -> int:
    x = max(0, min(a.right, b.right) - max(a.left, b.left))
    y = max(0, min(a.bottom, b.bottom) - max(a.top, b.top))
    return x * y


def box_contains(outer: Box, inner: Box, tolerance: int = 0) -> bool:
    return (
        outer.left - tolerance <= inner.left
        and outer.top - tolerance <= inner.top
        and outer.right + tolerance >= inner.right
        and outer.bottom + tolerance >= inner.bottom
    )


def point_in_box(x: int, y: int, box: Box, tolerance: int = 0) -> bool:
    return (
        box.left - tolerance <= x <= box.right + tolerance
        and box.top - tolerance <= y <= box.bottom + tolerance
    )


def line_intersects_box_interior(line: LineShape, box: Box, tolerance: int) -> bool:
    expanded = Box(box.left - tolerance, box.top - tolerance, box.width + 2 * tolerance, box.height + 2 * tolerance)
    if max(line.x1, line.x2) < expanded.left or min(line.x1, line.x2) > expanded.right:
        return False
    if max(line.y1, line.y2) < expanded.top or min(line.y1, line.y2) > expanded.bottom:
        return False
    if point_in_box(line.x1, line.y1, expanded, 0) or point_in_box(line.x2, line.y2, expanded, 0):
        return False
    if line.x1 == line.x2:
        return expanded.left < line.x1 < expanded.right
    if line.y1 == line.y2:
        return expanded.top < line.y1 < expanded.bottom

    # Check whether the segment crosses any side of the box.
    dx = line.x2 - line.x1
    dy = line.y2 - line.y1
    for x in (expanded.left, expanded.right):
        t = (x - line.x1) / dx
        if 0 < t < 1:
            y = line.y1 + t * dy
            if expanded.top < y < expanded.bottom:
                return True
    for y in (expanded.top, expanded.bottom):
        t = (y - line.y1) / dy
        if 0 < t < 1:
            x = line.x1 + t * dx
            if expanded.left < x < expanded.right:
                return True
    return False


def center_delta_in(a: Box, b: Box) -> float:
    return abs(a.center_y - b.center_y) / EMU_PER_INCH


def center_xy_delta_in(a: Box, b: Box) -> tuple[float, float]:
    return (
        abs(a.center_x - b.center_x) / EMU_PER_INCH,
        abs(a.center_y - b.center_y) / EMU_PER_INCH,
    )


def normalize_rotation(angle: float) -> float:
    while angle > 180:
        angle -= 360
    while angle <= -180:
        angle += 360
    return angle


def point_to_segment_distance_in(px: float, py: float, line: LineShape) -> float:
    vx = line.x2 - line.x1
    vy = line.y2 - line.y1
    wx = px - line.x1
    wy = py - line.y1
    length2 = vx * vx + vy * vy
    if length2 <= 0:
        return math.hypot(px - line.x1, py - line.y1) / EMU_PER_INCH
    t = max(0.0, min(1.0, (wx * vx + wy * vy) / length2))
    proj_x = line.x1 + t * vx
    proj_y = line.y1 + t * vy
    return math.hypot(px - proj_x, py - proj_y) / EMU_PER_INCH


def in_to_emu(value: float) -> int:
    return int(value * EMU_PER_INCH)


def text_units(text: str) -> float:
    units = 0.0
    for ch in text:
        if ch.isspace():
            units += 0.35
        elif "\u4e00" <= ch <= "\u9fff":
            units += 1.0
        elif ch in "，。；：、,.·:/+-":
            units += 0.45
        else:
            units += 0.58
    return units


def estimate_wrapped_line_count(raw_text: str, font_pt: float, box_width_pt: float) -> int:
    usable_width_pt = max(1.0, box_width_pt * 0.90)
    total = 0
    for source_line in raw_text.splitlines() or [raw_text]:
        line = source_line.strip()
        if not line:
            total += 1
            continue
        estimated_width_pt = text_units(line) * font_pt
        total += max(1, math.ceil(estimated_width_pt / usable_width_pt))
    return total


def has_cjk(text: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


def mixed_language_terms(text: str) -> list[str]:
    if not has_cjk(text):
        return []

    terms: list[str] = []
    for word in re.findall(r"[A-Za-z][A-Za-z0-9_&+/\-]{2,}", text):
        if "_" in word:
            continue
        parts = [part for part in re.split(r"[&+/\-]", word) if part]
        if parts and all(part.upper() in ALLOWED_MIXED_TERMS for part in parts):
            continue
        if word.upper() in ALLOWED_MIXED_TERMS:
            continue
        if word not in terms:
            terms.append(word)
    return terms


def is_compact_row_label(text: str, box: Box) -> bool:
    clean = text.strip()
    if not clean or clean.lower().startswith(("source:", "note:", "decision ")):
        return False
    if box.width > 2.7 * EMU_PER_INCH or box.height > 0.78 * EMU_PER_INCH:
        return False
    return text_units(clean) <= 32.0


def is_oval_shape(shape, oval_type) -> bool:
    try:
        return shape.auto_shape_type == oval_type
    except Exception:
        return False


def shape_solid_rgb(shape) -> tuple[int, int, int] | None:
    try:
        rgb = shape.fill.fore_color.rgb
    except Exception:
        return None
    if rgb is None:
        return None
    value = str(rgb)
    if len(value) != 6:
        return None
    try:
        return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
    except ValueError:
        return None


def text_run_rgb(shape) -> tuple[int, int, int] | None:
    try:
        paragraphs = shape.text_frame.paragraphs
    except Exception:
        return None
    for paragraph in paragraphs:
        for run in paragraph.runs:
            try:
                rgb = run.font.color.rgb
            except Exception:
                rgb = None
            if rgb is None:
                continue
            value = str(rgb)
            if len(value) != 6:
                continue
            try:
                return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
            except ValueError:
                continue
    return None


def brightness(rgb: tuple[int, int, int]) -> float:
    red, green, blue = rgb
    return red * 0.299 + green * 0.587 + blue * 0.114


def relative_luminance(rgb: tuple[int, int, int]) -> float:
    values = []
    for channel in rgb:
        normalized = channel / 255.0
        if normalized <= 0.03928:
            values.append(normalized / 12.92)
        else:
            values.append(((normalized + 0.055) / 1.055) ** 2.4)
    return 0.2126 * values[0] + 0.7152 * values[1] + 0.0722 * values[2]


def contrast_ratio(foreground: tuple[int, int, int], background: tuple[int, int, int]) -> float:
    fg = relative_luminance(foreground)
    bg = relative_luminance(background)
    lighter = max(fg, bg)
    darker = min(fg, bg)
    return (lighter + 0.05) / (darker + 0.05)


def color_hex(rgb: tuple[int, int, int]) -> str:
    return "#{:02X}{:02X}{:02X}".format(*rgb)


def is_punctuation_or_orphan_line(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if re.fullmatch(r"[。；，、,.!?！？:：;]+", stripped):
        return True
    if len(stripped) <= 1 and has_cjk(stripped):
        return True
    return False


def likely_short_wrapped_tail(raw_text: str, font_pt: float, box_width_pt: float) -> bool:
    usable_units = max(1.0, (box_width_pt * 0.90) / max(font_pt, 1.0))
    for source_line in raw_text.splitlines() or [raw_text]:
        stripped = source_line.strip()
        if not stripped:
            continue
        units = text_units(stripped)
        estimated_lines = math.ceil(units / usable_units)
        if estimated_lines < 2:
            continue
        last_line_units = units - usable_units * (estimated_lines - 1)
        if 0 < last_line_units <= ORPHAN_LINE_UNITS:
            return True
    return False


def choose_main_title(text_shapes: list[TextShape], slide_w: int, slide_h: int) -> TextShape | None:
    candidates = [
        item
        for item in text_shapes
        if item.max_font_pt >= 20.0
        and item.box.top < slide_h * 0.35
        and item.box.left < slide_w * 0.45
        and len(item.text.strip()) >= 4
    ]
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: (item.box.top, item.box.left, -item.max_font_pt))[0]


def choose_header_label(text_shapes: list[TextShape], title: TextShape | None, slide_h: int) -> TextShape | None:
    title_label = title.label if title else None
    candidates = [
        item
        for item in text_shapes
        if item.label != title_label
        and item.box.top < min(in_to_emu(0.75), int(slide_h * 0.14))
        and item.max_font_pt <= 14.0
        and len(item.text.strip()) <= 40
    ]
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: (item.box.top, item.box.left))[0]


def choose_header_rule(horizontal_rules: list[tuple[str, Box]], slide_h: int) -> tuple[str, Box] | None:
    candidates = [
        (label, box)
        for label, box in horizontal_rules
        if box.top < min(in_to_emu(0.75), int(slide_h * 0.14))
        and box.width >= in_to_emu(0.6)
        and box.width <= in_to_emu(3.0)
    ]
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: (item[1].top, item[1].left))[0]


def add_header_alignment_findings(metrics: list[SlideHeaderMetrics], findings: list[str]) -> None:
    usable = [item for item in metrics if item.title_box is not None]
    if len(usable) < 2:
        return

    title_lefts = [item.title_box.left for item in usable if item.title_box is not None]
    title_tops = [item.title_box.top for item in usable if item.title_box is not None]
    title_left_drift = (max(title_lefts) - min(title_lefts)) / EMU_PER_INCH
    title_top_drift = (max(title_tops) - min(title_tops)) / EMU_PER_INCH
    if title_left_drift > TITLE_DRIFT_TOLERANCE_IN or title_top_drift > TITLE_DRIFT_TOLERANCE_IN:
        findings.append(
            "- warning: deck-level title alignment drift: main title positions vary "
            f"by {title_left_drift:.2f} in horizontally and {title_top_drift:.2f} in vertically; "
            "generated decks should keep repeated title baselines locked unless a slide type intentionally changes."
        )

    rule_items = [item for item in usable if item.rule_box is not None and item.title_box is not None]
    if len(rule_items) >= 2:
        rule_lefts = [item.rule_box.left for item in rule_items if item.rule_box is not None]
        rule_tops = [item.rule_box.top for item in rule_items if item.rule_box is not None]
        rule_left_drift = (max(rule_lefts) - min(rule_lefts)) / EMU_PER_INCH
        rule_top_drift = (max(rule_tops) - min(rule_tops)) / EMU_PER_INCH
        if rule_left_drift > TITLE_DRIFT_TOLERANCE_IN or rule_top_drift > TITLE_DRIFT_TOLERANCE_IN:
            findings.append(
                "- warning: deck-level header accent drift: top accent rule positions vary "
                f"by {rule_left_drift:.2f} in horizontally and {rule_top_drift:.2f} in vertically."
            )

    floating_rule_slides: list[tuple[int, str | None, str | None, float]] = []
    detached_header_slides: list[tuple[int, str | None, float, float]] = []
    for item in usable:
        title = item.title_box
        rule = item.rule_box
        header = item.header_box
        if title is None:
            continue
        if rule is not None:
            rule_left_delta = abs(rule.left - title.left) / EMU_PER_INCH
            rule_right_delta = abs(rule.right - title.left) / EMU_PER_INCH
            if min(rule_left_delta, rule_right_delta) > HEADER_RELATION_TOLERANCE_IN:
                floating_rule_slides.append((item.slide_idx, item.rule_label, item.title_label, min(rule_left_delta, rule_right_delta)))
        if header is not None:
            header_title_delta = abs(header.left - title.left) / EMU_PER_INCH
            header_rule_delta = abs(header.left - rule.right) / EMU_PER_INCH if rule is not None else 0.0
            if header_title_delta > HEADER_RELATION_TOLERANCE_IN and header_rule_delta > HEADER_RELATION_TOLERANCE_IN:
                detached_header_slides.append((item.slide_idx, item.header_label, header_title_delta, header_rule_delta))

    if len(floating_rule_slides) >= 3:
        slide_list = ", ".join(str(item[0]) for item in floating_rule_slides[:12])
        max_delta = max(item[3] for item in floating_rule_slides)
        findings.append(
            "- warning: deck-level top header grid issue: top accent rules are not tied to the main title left edge "
            f"on slides {slide_list} (max nearest delta {max_delta:.2f} in). "
            "This often appears as a floating horizontal line above every title; fix the master/header component before slide-level polish."
        )
    else:
        for slide_idx, rule_label, title_label, delta in floating_rule_slides:
            findings.append(
                f"- warning: slide {slide_idx}: top accent rule `{rule_label}` is not tied to the main title left edge "
                f"`{title_label}` (nearest delta {delta:.2f} in); check whether the header line is visually floating instead of belonging to the title grid."
            )

    if len(detached_header_slides) >= 3:
        slide_list = ", ".join(str(item[0]) for item in detached_header_slides[:12])
        findings.append(
            "- warning: deck-level top header label issue: section labels are detached from both the title grid and accent rule "
            f"on slides {slide_list}; fix the repeated header component."
        )
    else:
        for slide_idx, header_label, header_title_delta, header_rule_delta in detached_header_slides:
            findings.append(
                f"- warning: slide {slide_idx}: top header label `{header_label}` is not aligned with the main title "
                f"or adjacent accent rule (title delta {header_title_delta:.2f} in, rule delta {header_rule_delta:.2f} in); check top-title/header grid."
            )


def export_with_powerpoint(pptx: Path, export_dir: Path) -> str:
    try:
        import win32com.client  # type: ignore
        import pythoncom  # type: ignore
    except Exception as exc:  # pragma: no cover - depends on local Windows setup
        return export_with_powershell(pptx, export_dir, f"pywin32 unavailable ({exc})")

    app = None
    pres = None
    exported_count = 0
    export_error = None
    quit_warning = None
    try:  # pragma: no cover - depends on local Windows setup
        pythoncom.CoInitialize()
        app = win32com.client.Dispatch("PowerPoint.Application")
        pres = app.Presentations.Open(str(pptx), WithWindow=False)
        if export_dir.exists():
            shutil.rmtree(export_dir)
        export_dir.mkdir(parents=True, exist_ok=True)
        for idx, slide in enumerate(pres.Slides, start=1):
            slide.Export(str(export_dir / f"slide_{idx:02d}.png"), "PNG", 1920, 1080)
            exported_count += 1
    except Exception as exc:
        export_error = exc
    finally:
        if pres is not None:
            try:
                pres.Close()
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception as exc:
                quit_warning = str(exc)
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

    if export_error is not None:
        return export_with_powershell(pptx, export_dir, f"pywin32 export failed ({export_error})")

    png_count = len(list(export_dir.glob("slide_*.png")))
    if png_count != exported_count:
        return f"PowerPoint export failed: expected {exported_count} PNG files, found {png_count} in `{export_dir}`."
    if quit_warning:
        return f"Exported slide images to `{export_dir}`. PowerPoint quit warning: {quit_warning}"
    return f"Exported slide images to `{export_dir}`."


def export_with_powershell(pptx: Path, export_dir: Path, reason: str) -> str:
    if sys.platform != "win32":
        return f"PowerPoint export skipped: {reason}."

    script = r"""
param(
    [Parameter(Mandatory=$true)][string]$Pptx,
    [Parameter(Mandatory=$true)][string]$ExportDir
)
$ErrorActionPreference = "Stop"
$app = $null
$pres = $null
if (Test-Path -LiteralPath $ExportDir) {
    Remove-Item -LiteralPath $ExportDir -Recurse -Force
}
New-Item -ItemType Directory -Path $ExportDir -Force | Out-Null
try {
    $app = New-Object -ComObject PowerPoint.Application
    $pres = $app.Presentations.Open($Pptx, $true, $false, $false)
    $count = 0
    foreach ($slide in $pres.Slides) {
        $count += 1
        $png = Join-Path $ExportDir ("slide_{0:D2}.png" -f $count)
        $slide.Export($png, "PNG", 1920, 1080)
    }
    Write-Output ("exported_count={0}" -f $count)
}
finally {
    if ($pres -ne $null) {
        try { $pres.Close() } catch {}
    }
    if ($app -ne $null) {
        try { $app.Quit() } catch {}
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""
    with tempfile.TemporaryDirectory() as tmp:
        ps1 = Path(tmp) / "export_ppt.ps1"
        ps1.write_text(script, encoding="utf-8")
        proc = subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(ps1),
                "-Pptx",
                str(pptx),
                "-ExportDir",
                str(export_dir),
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )

    png_count = len(list(export_dir.glob("slide_*.png")))
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout).strip().splitlines()
        message = detail[-1] if detail else "unknown PowerShell COM error"
        return f"PowerPoint export failed: {message}"
    if png_count == 0:
        return f"PowerPoint export failed: PowerShell COM produced no PNG files in `{export_dir}`."
    return f"Exported slide images to `{export_dir}` via PowerShell COM fallback ({reason})."


def box_to_inches(box: Box | None) -> list[str]:
    if box is None:
        return ["", "", "", ""]
    return [
        f"{box.left / EMU_PER_INCH:.3f}",
        f"{box.top / EMU_PER_INCH:.3f}",
        f"{box.width / EMU_PER_INCH:.3f}",
        f"{box.height / EMU_PER_INCH:.3f}",
    ]


def summarize_layout_audit(metrics: list[SlideHeaderMetrics]) -> list[str]:
    usable = [item for item in metrics if item.title_box is not None]
    if not usable:
        return ["- No repeated title/header metrics were detected; rely on PNG review and manual layout contract."]

    title_lefts = [item.title_box.left for item in usable if item.title_box is not None]
    title_tops = [item.title_box.top for item in usable if item.title_box is not None]
    title_left_drift = (max(title_lefts) - min(title_lefts)) / EMU_PER_INCH
    title_top_drift = (max(title_tops) - min(title_tops)) / EMU_PER_INCH
    rule_count = len([item for item in usable if item.rule_box is not None])
    header_count = len([item for item in usable if item.header_box is not None])
    return [
        f"- Title metrics detected on {len(usable)} slide(s); left drift {title_left_drift:.2f} in, top drift {title_top_drift:.2f} in.",
        f"- Header labels detected on {header_count} slide(s); top accent rules detected on {rule_count} slide(s).",
        "- Use `layout_audit.csv` to compare title/header/rule positions across slides and after regeneration.",
    ]


def write_layout_audit(path: Path, metrics: list[SlideHeaderMetrics]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "slide",
                "title_text",
                "title_left_in",
                "title_top_in",
                "title_width_in",
                "title_height_in",
                "header_text",
                "header_left_in",
                "header_top_in",
                "header_width_in",
                "header_height_in",
                "rule_left_in",
                "rule_top_in",
                "rule_width_in",
                "rule_height_in",
            ]
        )
        for item in metrics:
            writer.writerow(
                [
                    item.slide_idx,
                    item.title_text,
                    *box_to_inches(item.title_box),
                    item.header_text,
                    *box_to_inches(item.header_box),
                    *box_to_inches(item.rule_box),
                ]
            )


def collect_findings(pptx: Path) -> tuple[list[str], str, list[SlideHeaderMetrics], list[str]]:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.text import MSO_ANCHOR  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
    except Exception as exc:
        return ([f"- blocker: unable to import python-pptx: {exc}"], "unknown", [], [])

    prs = Presentation(str(pptx))
    slide_w = prs.slide_width
    slide_h = prs.slide_height
    margin = int(EDGE_MARGIN_IN * EMU_PER_INCH)
    findings: list[str] = []
    header_metrics: list[SlideHeaderMetrics] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        text_boxes: list[tuple[str, Box]] = []
        text_shapes: list[TextShape] = []
        horizontal_rules: list[tuple[str, Box]] = []
        line_shapes: list[LineShape] = []
        visual_shapes: list[VisualShape] = []
        oval_shapes: list[tuple[str, Box]] = []
        filled_shapes: list[tuple[str, Box, tuple[int, int, int]]] = []
        dark_canvas = False
        large_light_panels: list[tuple[str, Box, float]] = []
        large_mid_panels: list[tuple[str, Box]] = []
        thin_mid_rules: list[tuple[str, Box]] = []
        short_label_rules: list[tuple[str, Box]] = []
        for shape_label, shape, shape_idx in iter_slide_shapes(slide):
            if not all(hasattr(shape, attr) for attr in ("left", "top", "width", "height")):
                continue
            box = Box(int(shape.left), int(shape.top), int(shape.width), int(shape.height))
            label = f"slide {slide_idx}, {shape_label}"
            fill_rgb = shape_solid_rgb(shape)
            if fill_rgb is not None:
                if box.width * box.height >= in_to_emu(0.08) * in_to_emu(0.08):
                    filled_shapes.append((label, box, fill_rgb))
                fill_brightness = brightness(fill_rgb)
                if (
                    box.left <= margin
                    and box.top <= margin
                    and box.width >= slide_w * 0.88
                    and box.height >= slide_h * 0.88
                    and fill_brightness < 70
                ):
                    dark_canvas = True
                if (
                    box.width >= in_to_emu(5.0)
                    and box.height >= in_to_emu(0.75)
                    and fill_brightness > 225
                    and box.width * box.height >= slide_w * slide_h * 0.08
                ):
                    large_light_panels.append((label, box, fill_brightness))
                if (
                    box.width >= in_to_emu(6.0)
                    and in_to_emu(0.75) <= box.height <= in_to_emu(2.0)
                    and in_to_emu(1.0) <= box.top <= in_to_emu(5.3)
                    and fill_brightness <= 170
                ):
                    large_mid_panels.append((label, box))

            if is_oval_shape(shape, MSO_SHAPE.OVAL):
                oval_shapes.append((label, box))

            try:
                shape_type = shape.shape_type
            except Exception:
                shape_type = None
            is_picture = shape_type == MSO_SHAPE_TYPE.PICTURE
            is_text_box = shape_type == MSO_SHAPE_TYPE.TEXT_BOX
            is_line_shape = shape_type == MSO_SHAPE_TYPE.LINE
            if is_line_shape and all(hasattr(shape, attr) for attr in ("begin_x", "begin_y", "end_x", "end_y")):
                x1 = int(shape.begin_x)
                y1 = int(shape.begin_y)
                x2 = int(shape.end_x)
                y2 = int(shape.end_y)
                try:
                    width_pt = float(shape.line.width.pt)
                except Exception:
                    width_pt = 1.0
                if math.hypot(x2 - x1, y2 - y1) >= in_to_emu(0.25):
                    line_shapes.append(LineShape(label, x1, y1, x2, y2, width_pt))
            try:
                is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            except Exception:
                is_picture = False
            try:
                is_text_box = shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
            except Exception:
                is_text_box = False
            if is_picture:
                pic_w = box.width / EMU_PER_INCH
                pic_h = box.height / EMU_PER_INCH
                pic_area = pic_w * pic_h
                if (
                    pic_area >= 0.35
                    and (
                        pic_w < MIN_EVIDENCE_IMAGE_WIDTH_IN
                        or pic_h < MIN_EVIDENCE_IMAGE_HEIGHT_IN
                        or pic_area < MIN_EVIDENCE_IMAGE_AREA_IN2
                    )
                ):
                    findings.append(
                        f"- warning: {label}: picture is only {pic_w:.1f} x {pic_h:.1f} in; "
                        "if it contains labels, code, terminal text, or a UI screenshot, it is likely unreadable at presentation size. "
                        "Crop/enlarge it or replace it with a simplified diagram."
                    )

            if (
                (fill_rgb is not None or is_picture or is_oval_shape(shape, MSO_SHAPE.OVAL))
                and not is_text_box
                and box.width * box.height >= in_to_emu(0.06) * in_to_emu(0.06)
                and not (box.width >= slide_w * 0.88 and box.height >= slide_h * 0.88)
            ):
                visual_shapes.append(VisualShape(label, box, "picture" if is_picture else "shape"))

            is_long_separator = box.width > 3.0 * EMU_PER_INCH and box.height < 0.07 * EMU_PER_INCH
            is_top_accent_rule = (
                box.top < min(in_to_emu(0.75), int(slide_h * 0.14))
                and box.width > 0.6 * EMU_PER_INCH
                and box.height < 0.08 * EMU_PER_INCH
            )
            if is_long_separator or is_top_accent_rule:
                horizontal_rules.append((label, box))
            if (
                in_to_emu(0.25) <= box.width <= in_to_emu(2.4)
                and box.height < in_to_emu(0.07)
                and in_to_emu(1.0) <= box.top <= in_to_emu(5.8)
            ):
                thin_mid_rules.append((label, box))
            if (
                in_to_emu(0.25) <= box.width <= in_to_emu(1.25)
                and box.height < in_to_emu(0.075)
                and in_to_emu(0.20) <= box.top <= in_to_emu(6.75)
            ):
                short_label_rules.append((label, box))

            if box.left < -margin or box.top < -margin or box.right > slide_w + margin or box.bottom > slide_h + margin:
                findings.append(f"- warning: {label}: object appears outside slide bounds.")
            elif box.left < margin or box.top < margin or slide_w - box.right < margin or slide_h - box.bottom < margin:
                findings.append(f"- info: {label}: object is close to slide edge.")

            if getattr(shape, "has_text_frame", False):
                raw_text = (shape.text or "").strip()
                text = raw_text.replace("\n", " ")
                if text:
                    vertical_middle = shape.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
                    text_boxes.append((label, box))
                    sizes: list[float] = []
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size is not None:
                                sizes.append(run.font.size.pt)
                    min_font_pt = min(sizes) if sizes else 12.0
                    max_font_pt = max(sizes) if sizes else min_font_pt
                    text_color = text_run_rgb(shape)
                    if sizes and min(sizes) < MIN_VISIBLE_FONT_PT:
                        findings.append(
                            f"- blocker: {label}: visible text is below 9 pt ({min(sizes):.1f} pt). "
                            "Generated PPT previews and component-library samples must keep required visible text at 9 pt or larger. "
                            f"Text starts: `{text[:70]}`."
                        )
                    elif sizes and min(sizes) < MIN_BODY_FONT_PT:
                        findings.append(
                            f"- warning: {label}: small text detected ({min(sizes):.1f} pt). Text starts: `{text[:70]}`."
                        )
                    rotation = normalize_rotation(float(getattr(shape, "rotation", 0.0) or 0.0))
                    text_shapes.append(TextShape(label, box, text, vertical_middle, min_font_pt, max_font_pt, text_color, rotation))
                    font_pt = min_font_pt
                    box_width_pt = max(1.0, box.width / EMU_PER_INCH * 72)
                    box_height_pt = max(1.0, box.height / EMU_PER_INCH * 72)
                    estimated_lines = estimate_wrapped_line_count(raw_text, font_pt, box_width_pt)
                    estimated_text_height_pt = estimated_lines * font_pt * 1.18
                    if font_pt < 20.0 and estimated_lines > 1 and estimated_text_height_pt > box_height_pt * 1.18:
                        findings.append(
                            f"- warning: {label}: likely text overflow or clipped descenders; estimated {estimated_lines} wrapped line(s) at {font_pt:.1f} pt in a {box_height_pt:.1f} pt-high box. Text starts: `{text[:70]}`."
                        )
                    for line in raw_text.splitlines() or [raw_text]:
                        if not line.strip():
                            continue
                        estimated_width_pt = text_units(line.strip()) * font_pt
                        if estimated_width_pt > box_width_pt * 0.92:
                            findings.append(
                                f"- warning: {label}: long unwrapped text may need a manual line break. Text starts: `{line.strip()[:70]}`."
                            )
                            break
                    if likely_short_wrapped_tail(raw_text, font_pt, box_width_pt):
                        findings.append(
                            f"- warning: {label}: wrapped text likely leaves a very short final line or orphan punctuation; shorten text, widen the box, or add an intentional line break. Text starts: `{text[:70]}`."
                        )
                    for line in raw_text.splitlines() or [raw_text]:
                        if is_punctuation_or_orphan_line(line):
                            findings.append(
                                f"- warning: {label}: visible orphan line `{line.strip()}`; do not leave punctuation or a one-character CJK tail alone in a polished deck."
                            )
                            break
                    try:
                        margins = [
                            shape.text_frame.margin_left / EMU_PER_INCH,
                            shape.text_frame.margin_right / EMU_PER_INCH,
                            shape.text_frame.margin_top / EMU_PER_INCH,
                            shape.text_frame.margin_bottom / EMU_PER_INCH,
                        ]
                    except Exception:
                        margins = []
                    if (not is_text_box) and fill_rgb is not None and margins and min(margins) < MIN_SHAPE_TEXT_MARGIN_IN and box.width > in_to_emu(0.6) and box.height > in_to_emu(0.25):
                        severity = "blocker" if min(margins) < 0.08 else "warning"
                        findings.append(
                            f"- {severity}: {label}: text frame has small internal margin ({min(margins):.2f} in); "
                            "text may look too close to the shape border. Use shape-owned padding or a helper-owned safe inset."
                        )
                    if (
                        slide_idx > 1
                        and shape_idx <= 5
                        and box.left < slide_w * 0.12
                        and box.right > slide_w * 0.46
                        and slide_h * 0.45 < box.top < slide_h * 0.75
                        and len(text) > 24
                    ):
                        findings.append(
                            f"- warning: {label}: section-style subtitle may press into the template art area; check exported PNG."
                        )
                    for pattern, message in COMMON_VISIBLE_TEXT_RISKS:
                        if pattern.search(text):
                            if pattern.pattern == r"\bdemo\b" and "MPPT Demo" in text:
                                continue
                            findings.append(f"- warning: {label}: {message}. Text starts: `{text[:70]}`.")
                            break
                    mixed_terms = mixed_language_terms(text)
                    if mixed_terms:
                        findings.append(
                            f"- info: {label}: mixed Chinese/English terms may need audience review ({', '.join(mixed_terms[:5])}). Text starts: `{text[:70]}`."
                        )

        for item in text_shapes:
            candidate_panels: list[tuple[int, str, Box, tuple[int, int, int]]] = []
            for panel_label, panel_box, panel_rgb in filled_shapes:
                if panel_label == item.label:
                    continue
                if panel_box.width > slide_w * 0.88 and panel_box.height > slide_h * 0.88:
                    continue
                if (
                    panel_box.left <= item.box.center_x <= panel_box.right
                    and panel_box.top <= item.box.center_y <= panel_box.bottom
                    and panel_box.width * panel_box.height > item.box.width * item.box.height
                ):
                    candidate_panels.append((panel_box.width * panel_box.height, panel_label, panel_box, panel_rgb))
            if not candidate_panels:
                continue
            _, panel_label, panel_box, panel_rgb = sorted(candidate_panels, key=lambda value: value[0])[0]
            gaps = {
                "left": (item.box.left - panel_box.left) / EMU_PER_INCH,
                "right": (panel_box.right - item.box.right) / EMU_PER_INCH,
                "top": (item.box.top - panel_box.top) / EMU_PER_INCH,
                "bottom": (panel_box.bottom - item.box.bottom) / EMU_PER_INCH,
            }
            min_edge, min_gap = min(gaps.items(), key=lambda pair: pair[1])
            if min_gap < -0.02:
                findings.append(
                    f"- blocker: {item.label}: text box extends outside containing frame {panel_label} on the {min_edge} edge ({min_gap:.2f} in). "
                    "Generated framed text must reserve safe bounds before rendering."
                )
            elif min_gap < MIN_TEXT_CONTAINER_INSET_IN and item.max_font_pt >= 8.0:
                findings.append(
                    f"- warning: {item.label}: text box is only {min_gap:.2f} in from the {min_edge} edge of containing frame {panel_label}; "
                    f"use at least {MIN_TEXT_CONTAINER_INSET_IN:.2f} in helper-owned padding or resize/split the content."
                )
            if item.color_rgb is not None:
                ratio = contrast_ratio(item.color_rgb, panel_rgb)
                required = MIN_LARGE_TEXT_CONTRAST if item.max_font_pt >= 18.0 else MIN_NORMAL_TEXT_CONTRAST
                if ratio < required:
                    findings.append(
                        f"- blocker: {item.label}: low text/background contrast against containing frame {panel_label} "
                        f"(ratio {ratio:.2f}:1, text {color_hex(item.color_rgb)}, background {color_hex(panel_rgb)}, required about {required:.1f}:1). "
                        "Fix foreground/background colors, add a style-matched surface, or change typography before handoff."
                    )

        if dark_canvas:
            for panel_label, panel_box, fill_brightness in large_light_panels:
                findings.append(
                    f"- warning: {panel_label}: large light panel on a dark visual system "
                    f"({panel_box.width / EMU_PER_INCH:.1f} x {panel_box.height / EMU_PER_INCH:.1f} in, brightness {fill_brightness:.0f}). "
                    "Do not fix readability by washing a region with an unrelated background; use the style palette, typography, and component surface contrast instead."
                )

        for panel_label, panel_box in large_mid_panels:
            contained_rules = [
                rule_label
                for rule_label, rule_box in thin_mid_rules
                if panel_box.left <= rule_box.center_x <= panel_box.right
                and panel_box.top <= rule_box.center_y <= panel_box.bottom
            ]
            contained_texts = [
                item.label
                for item in text_shapes
                if panel_box.left <= item.box.center_x <= panel_box.right
                and panel_box.top <= item.box.center_y <= panel_box.bottom
            ]
            if len(contained_rules) >= 2 and len(contained_texts) >= 4:
                findings.append(
                    f"- warning: {panel_label}: likely shared process band contains {len(contained_rules)} connector/rule segment(s) and {len(contained_texts)} text block(s). "
                    "Connector lines can look like they float inside one large container; use helper-owned step cards/nodes and connect card boundaries or centers instead."
                )

        title_shape = choose_main_title(text_shapes, slide_w, slide_h)
        header_shape = choose_header_label(text_shapes, title_shape, slide_h)
        rule_item = choose_header_rule(horizontal_rules, slide_h)
        header_metrics.append(
            SlideHeaderMetrics(
                slide_idx=slide_idx,
                title_label=title_shape.label if title_shape else None,
                title_box=title_shape.box if title_shape else None,
                title_text=title_shape.text if title_shape else "",
                header_label=header_shape.label if header_shape else None,
                header_box=header_shape.box if header_shape else None,
                header_text=header_shape.text if header_shape else "",
                rule_label=rule_item[0] if rule_item else None,
                rule_box=rule_item[1] if rule_item else None,
            )
        )

        for i, (label_a, box_a) in enumerate(text_boxes):
            for label_b, box_b in text_boxes[i + 1 :]:
                overlap = box_overlap(box_a, box_b)
                if overlap == 0:
                    continue
                smaller_area = max(1, min(box_a.width * box_a.height, box_b.width * box_b.height))
                if overlap / smaller_area > 0.12:
                    findings.append(f"- warning: {label_a} overlaps text box at {label_b}.")

        for idx, shape_a in enumerate(visual_shapes):
            for shape_b in visual_shapes[idx + 1 :]:
                overlap = box_overlap(shape_a.box, shape_b.box)
                if overlap == 0:
                    continue
                smaller_area = max(1, min(shape_a.box.width * shape_a.box.height, shape_b.box.width * shape_b.box.height))
                overlap_ratio = overlap / smaller_area
                overlap_area_in2 = overlap / (EMU_PER_INCH * EMU_PER_INCH)
                tolerance = in_to_emu(0.01)
                if (
                    (box_contains(shape_a.box, shape_b.box, tolerance) or box_contains(shape_b.box, shape_a.box, tolerance))
                    and overlap_ratio > 0.92
                ):
                    continue
                if overlap_area_in2 >= MIN_GRAPHIC_OVERLAP_AREA_IN2 and overlap_ratio > GRAPHIC_OVERLAP_RATIO:
                    findings.append(
                        f"- warning: {shape_a.label} graphically overlaps {shape_b.label} "
                        f"(overlap {overlap_area_in2:.2f} in^2, {overlap_ratio:.0%} of smaller object). "
                        "Inspect exported PNG; icons, symbols, regions, and helper-owned shapes should not cover each other unless this is an intentional child component."
                    )

        for visual in visual_shapes:
            is_bar_like = (
                visual.box.width >= in_to_emu(0.55)
                and visual.box.height <= in_to_emu(0.32)
                and visual.box.width * visual.box.height >= in_to_emu(0.06) * in_to_emu(0.06)
            )
            if not is_bar_like:
                continue
            candidate_panels: list[tuple[int, str, Box]] = []
            for panel_label, panel_box, _panel_rgb in filled_shapes:
                if panel_label == visual.label:
                    continue
                if panel_box.width > slide_w * 0.88 and panel_box.height > slide_h * 0.88:
                    continue
                if panel_box.width * panel_box.height <= visual.box.width * visual.box.height * 3:
                    continue
                if panel_box.left <= visual.box.center_x <= panel_box.right and panel_box.top <= visual.box.center_y <= panel_box.bottom:
                    candidate_panels.append((panel_box.width * panel_box.height, panel_label, panel_box))
            if not candidate_panels:
                continue
            _, panel_label, panel_box = sorted(candidate_panels, key=lambda value: value[0])[0]
            gaps = {
                "left": (visual.box.left - panel_box.left) / EMU_PER_INCH,
                "right": (panel_box.right - visual.box.right) / EMU_PER_INCH,
                "top": (visual.box.top - panel_box.top) / EMU_PER_INCH,
                "bottom": (panel_box.bottom - visual.box.bottom) / EMU_PER_INCH,
            }
            min_edge, min_gap = min(gaps.items(), key=lambda pair: pair[1])
            if min_gap < -0.02:
                findings.append(
                    f"- blocker: {visual.label}: bar/graphic extends outside containing frame {panel_label} on the {min_edge} edge ({min_gap:.2f} in). "
                    "Size child bars from the parent slot and reserve visible insets before rendering."
                )
            elif min_gap < MIN_GRAPHIC_CONTAINER_INSET_IN:
                findings.append(
                    f"- warning: {visual.label}: bar/graphic is only {min_gap:.2f} in from the {min_edge} edge of containing frame {panel_label}; "
                    f"use at least {MIN_GRAPHIC_CONTAINER_INSET_IN:.2f} in parent-owned graphic inset."
                )

        connector_tolerance = in_to_emu(CONNECTOR_CROSSING_TOLERANCE_IN)
        for line_item in line_shapes:
            if line_item.width_pt < 0.6:
                continue
            for visual in visual_shapes:
                if line_intersects_box_interior(line_item, visual.box, connector_tolerance):
                    findings.append(
                        f"- warning: {line_item.label}: connector or rule may cross through {visual.label}; "
                        "line routes should attach at declared ports or stay in connector lanes instead of covering diagram objects."
                    )

        for item in text_shapes:
            if abs(item.rotation) < 8.0 or item.box.width < in_to_emu(0.30) or item.box.height < in_to_emu(0.08):
                continue
            for line_item in line_shapes:
                expanded = in_to_emu(0.18)
                if not (
                    min(line_item.x1, line_item.x2) - expanded <= item.box.center_x <= max(line_item.x1, line_item.x2) + expanded
                    and min(line_item.y1, line_item.y2) - expanded <= item.box.center_y <= max(line_item.y1, line_item.y2) + expanded
                ):
                    continue
                distance = point_to_segment_distance_in(item.box.center_x, item.box.center_y, line_item)
                if distance < 0.08:
                    findings.append(
                        f"- warning: {item.label}: rotated transition label may sit on diagonal connector {line_item.label} "
                        f"(center distance {distance:.2f} in). Offset the label parallel to the line so text does not cover the stroke."
                    )
                    break

        for rule_label, rule_box in horizontal_rules:
            for item in text_shapes:
                text_label, text_box, text = item.label, item.box, item.text
                if text_box.bottom <= rule_box.bottom:
                    continue
                gap = (text_box.top - rule_box.bottom) / EMU_PER_INCH
                horizontally_near = text_box.left < rule_box.right and text_box.right > rule_box.left
                if horizontally_near and 0 <= gap < SEPARATOR_TEXT_GAP_IN:
                    findings.append(
                        f"- warning: {text_label}: text is only {gap:.2f} in below horizontal separator {rule_label}; check visual air in exported PNG."
                    )

        for rule_label, rule_box in short_label_rules:
            candidates: list[tuple[int, TextShape]] = []
            for item in text_shapes:
                if not item.text.strip():
                    continue
                if re.fullmatch(r"\d{1,2}", item.text.strip()):
                    continue
                horizontal_gap = item.box.left - rule_box.right
                center_delta = abs(item.box.center_y - rule_box.center_y)
                if (
                    0 <= horizontal_gap <= in_to_emu(0.90)
                    and center_delta <= in_to_emu(0.35)
                    and item.box.height <= in_to_emu(0.50)
                ):
                    candidates.append((horizontal_gap + center_delta, item))
            if not candidates:
                continue
            item = sorted(candidates, key=lambda candidate: candidate[0])[0][1]
            delta = center_delta_in(rule_box, item.box)
            if delta > RULE_LABEL_CENTER_TOLERANCE_IN or not item.vertical_middle:
                findings.append(
                    f"- warning: {item.label}: short accent rule `{rule_label}` and adjacent label/callout text do not share a stable center line "
                    f"(center delta {delta:.2f} in; text middle anchor={item.vertical_middle}). "
                    "Use one rule-label helper or grouped callout component instead of independent offsets."
                )

        for item_a in text_shapes:
            label_a, box_a, text_a = item_a.label, item_a.box, item_a.text
            if not re.fullmatch(r"\d{1,2}", text_a.strip()):
                continue
            if box_a.width > 0.8 * EMU_PER_INCH or box_a.height > 0.5 * EMU_PER_INCH:
                continue
            circle_candidates: list[tuple[float, str, Box]] = []
            for oval_label, oval_box in oval_shapes:
                center_dx, center_dy = center_xy_delta_in(box_a, oval_box)
                if center_dx <= 0.10 and center_dy <= 0.10:
                    circle_candidates.append((center_dx + center_dy, oval_label, oval_box))
            if circle_candidates:
                _, oval_label, oval_box = sorted(circle_candidates, key=lambda item: item[0])[0]
                center_dx, center_dy = center_xy_delta_in(box_a, oval_box)
                same_box = (
                    abs(box_a.left - oval_box.left) <= in_to_emu(0.025)
                    and abs(box_a.top - oval_box.top) <= in_to_emu(0.025)
                    and abs(box_a.width - oval_box.width) <= in_to_emu(0.025)
                    and abs(box_a.height - oval_box.height) <= in_to_emu(0.025)
                )
                if center_dx > CIRCLE_TEXT_CENTER_TOLERANCE_IN or center_dy > CIRCLE_TEXT_CENTER_TOLERANCE_IN or not item_a.vertical_middle or not same_box:
                    findings.append(
                        f"- warning: {label_a}: numeric circle text is not owned by the circle component `{oval_label}` "
                        f"(center delta x/y {center_dx:.2f}/{center_dy:.2f} in; same-box={same_box}; middle anchor={item_a.vertical_middle}). "
                        "Use one circle badge helper so the text frame and oval share bounds."
                    )
                continue

            candidates: list[tuple[float, str, Box, bool]] = []
            for item_b in text_shapes:
                label_b, box_b, text_b, vertical_b = item_b.label, item_b.box, item_b.text, item_b.vertical_middle
                if label_a == label_b or re.fullmatch(r"\d{1,2}", text_b.strip()):
                    continue
                horizontal_gap = box_b.left - box_a.right
                if 0 <= horizontal_gap <= 2.6 * EMU_PER_INCH and abs(box_b.top - box_a.top) <= 0.75 * EMU_PER_INCH:
                    candidates.append((horizontal_gap + abs(box_b.top - box_a.top), label_b, box_b, vertical_b))
            if not candidates:
                continue
            row_items = sorted(candidates, key=lambda item: item[0])[:2]
            for _, label_b, box_b, vertical_b in row_items:
                delta = center_delta_in(box_a, box_b)
                if delta > BADGE_CENTER_TOLERANCE_IN:
                    findings.append(
                        f"- warning: {label_a} and adjacent label {label_b} may not be visually centered together (center delta {delta:.2f} in)."
                    )
                if delta > COMPONENT_CENTER_TOLERANCE_IN or not item_a.vertical_middle or not vertical_b:
                    findings.append(
                        f"- warning: {label_a} and adjacent row text {label_b} do not meet the generated-row component contract "
                        f"(center delta {delta:.2f} in; middle anchors={item_a.vertical_middle}/{vertical_b}). "
                        "For generated decks, number, title, and detail text should use one row helper with shared row height and middle anchors."
                    )
            if len(row_items) == 2:
                _, label_b, box_b, vertical_b = row_items[0]
                _, label_c, box_c, vertical_c = row_items[1]
                peer_delta = center_delta_in(box_b, box_c)
                if peer_delta > COMPONENT_CENTER_TOLERANCE_IN or not vertical_b or not vertical_c:
                    findings.append(
                        f"- warning: adjacent row texts {label_b} and {label_c} do not share one row center line "
                        f"(center delta {peer_delta:.2f} in; middle anchors={vertical_b}/{vertical_c}). "
                        "This is often invisible to geometry-only checks but obvious in rendered agenda rows."
                    )

        for oval_label, oval_box in oval_shapes:
            if oval_box.width > in_to_emu(0.45) or oval_box.height > in_to_emu(0.45):
                continue
            candidates: list[tuple[float, TextShape]] = []
            for item in text_shapes:
                if re.fullmatch(r"\d{1,2}", item.text.strip()):
                    continue
                horizontal_gap = item.box.left - oval_box.right
                center_delta = abs(item.box.center_y - oval_box.center_y)
                if 0 <= horizontal_gap <= in_to_emu(0.75) and center_delta <= in_to_emu(0.25):
                    candidates.append((horizontal_gap + center_delta, item))
            if not candidates:
                continue
            item = sorted(candidates, key=lambda candidate: candidate[0])[0][1]
            delta = center_delta_in(oval_box, item.box)
            if delta > ROW_LABEL_CENTER_TOLERANCE_IN or not item.vertical_middle:
                findings.append(
                    f"- warning: {oval_label} and adjacent note/list text {item.label} do not share one row center line "
                    f"(center delta {delta:.2f} in; text middle anchor={item.vertical_middle}). "
                    "Use one dot-label row helper or grouped component instead of separate marker and text offsets."
                )

        for item_a in text_shapes:
            label_a, box_a, text_a, vertical_a = item_a.label, item_a.box, item_a.text, item_a.vertical_middle
            if re.fullmatch(r"\d{1,2}", text_a.strip()) or not vertical_a or not is_compact_row_label(text_a, box_a):
                continue
            candidates: list[tuple[float, str, Box, bool]] = []
            for item_b in text_shapes:
                label_b, box_b, text_b, vertical_b = item_b.label, item_b.box, item_b.text, item_b.vertical_middle
                if label_a == label_b or not text_b.strip():
                    continue
                horizontal_gap = box_b.left - box_a.right
                if 0 <= horizontal_gap <= 3.4 * EMU_PER_INCH and abs(box_b.top - box_a.top) <= 0.6 * EMU_PER_INCH:
                    candidates.append((horizontal_gap + abs(box_b.top - box_a.top), label_b, box_b, vertical_b))
            if not candidates:
                continue
            for _, label_b, box_b, vertical_b in sorted(candidates, key=lambda item: item[0])[:2]:
                delta = center_delta_in(box_a, box_b)
                if delta > ROW_LABEL_CENTER_TOLERANCE_IN or not vertical_b:
                    findings.append(
                        f"- warning: {label_a} and adjacent row text {label_b} may not share a stable row center line (center delta {delta:.2f} in; middle anchors={vertical_a}/{vertical_b})."
                    )

    add_header_alignment_findings(header_metrics, findings)

    return findings, str(len(prs.slides)), header_metrics, summarize_layout_audit(header_metrics)


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

    parser = argparse.ArgumentParser()
    parser.add_argument("deck", type=Path)
    parser.add_argument("--out", type=Path, default=None, help="Report directory.")
    parser.add_argument("--no-export", action="store_true", help="Skip PowerPoint image export.")
    args = parser.parse_args()

    deck = args.deck.resolve()
    if not deck.exists():
        raise SystemExit(f"deck not found: {deck}")

    out_dir = args.out or (deck.parent / "visual_review" / deck.stem)
    out_dir.mkdir(parents=True, exist_ok=True)

    export_status = "Slide image export skipped by --no-export."
    if not args.no_export:
        export_status = export_with_powerpoint(deck, out_dir / "slides_png")

    findings, slide_count, header_metrics, layout_summary = collect_findings(deck)
    if not findings:
        findings = ["- info: no heuristic structural issues detected. Review exported images manually before release."]
    audit_path = out_dir / "layout_audit.csv"
    write_layout_audit(audit_path, header_metrics)

    report = out_dir / "visual_review_report.md"
    report.write_text(
        "\n".join(
            [
                f"# PPT Visual Review: {deck.name}",
                "",
                f"- Deck: `{deck}`",
                f"- Slides: {slide_count}",
                f"- Export: {export_status}",
                "",
                "## Findings",
                "",
                *findings,
                "",
                "## Layout Audit Evidence",
                "",
                f"- CSV: `{audit_path}`",
                *layout_summary,
                "",
                "## Manual Review Checklist",
                "",
                *[f"- {item}" for item in MANUAL_CHECKLIST],
            ]
        ),
        encoding="utf-8",
    )
    print(f"wrote {report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
